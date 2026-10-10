"""Export G06_HCI_CA3_Presentation.html to PDF and PPTX for the Drive submission.

Usage: python export_presentation.py            -> .pdf + .pptx
       python export_presentation.py --png       -> slide PNGs only (presentation_assets/slides/, for checking)
- PDF: Chrome "print to PDF" using the deck's print stylesheet (vector text, one 16:9 page per slide).
- PPTX: one full-slide 2x screenshot per slide (looks identical to the HTML; text is not editable —
  change build_presentation.py and re-export instead). Needs: pip install python-pptx
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from capture_screens import browser

HERE = Path(__file__).parent
DECK = HERE / "G06_HCI_CA3_Presentation.html"
SLIDES_DIR = HERE / "presentation_assets" / "slides"
PDF = DECK.with_suffix(".pdf")
PPTX = DECK.with_suffix(".pptx")


def screenshots(out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    n = len(re.findall(r'<section class="slide', DECK.read_text(encoding="utf-8")))
    paths = []
    for i in range(1, n + 1):
        png = out_dir / f"slide_{i:02d}.png"
        subprocess.run([browser(), "--headless=new", "--disable-gpu", "--hide-scrollbars",
                        "--force-device-scale-factor=2", "--window-size=1600,900", "--virtual-time-budget=3000",
                        f"--screenshot={png}", f"{DECK.resolve().as_uri()}?static#{i}"],
                       check=True, capture_output=True)
        paths.append(png)
    print(f"{n} slide screenshots")
    return paths


def pdf():
    subprocess.run([browser(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--virtual-time-budget=3000", f"--print-to-pdf={PDF}", DECK.resolve().as_uri()],
                   check=True, capture_output=True)
    print("wrote", PDF.name)


def pptx(paths):
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in (16:9)
    for p in paths:
        s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        s.shapes.add_picture(str(p), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(PPTX)
    print("wrote", PPTX.name)


if __name__ == "__main__":
    if "--png" in sys.argv:
        screenshots(SLIDES_DIR)
    else:
        tmp = Path(tempfile.mkdtemp())
        try:
            pdf()
            pptx(screenshots(tmp))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
