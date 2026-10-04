"""Build G04.docx: reuse the cover of ForgeFlow_Deployment_Report.docx, then append the report body.

Usage: python3 build_report.py      (needs implementation/results/summary.json and implementation/figures/*.png)
Citations are written as [@key] and numbered in order of first appearance (IEEE style).
"""
import copy
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
FIGS = HERE / "implementation" / "figures"
S = json.loads((HERE / "implementation" / "results" / "summary.json").read_text())
B, N = S["baseline"], S["nhi"]
US = S["gateway_us_per_call"]
TT = S["tamper_test"]

BLUE = RGBColor(0x1F, 0x38, 0x64)
BLACK = RGBColor(0, 0, 0)
RED = "B42D27"
FONT = "Times New Roman"

TITLE = ("Non-Human Identity and Scoped Delegation for Intrusion Detection and Forensic Attribution "
         "of AI-Orchestrated Attacks: A Case Study of the GTG-1002 Espionage Campaign (2025)")
MEMBERS = [("Raghav Sonchhatra", "23070122172"), ("Sanidhya Awasthi", "23070122192"),
           ("Faheemuddin Sayyed", "23070122196")]
PROBLEM = ("AI agents now call tools and delegate work to other agents, yet they usually run on broad, "
           "long-lived credentials that cannot be traced to an accountable person. The GTG-1002 campaign "
           "showed how this gap lets an AI-orchestrated intrusion move fast and stay unattributed. This study "
           "analyses the campaign and evaluates whether per-agent identities, scoped delegation and "
           "identity-aware logging improve prevention, intrusion detection and forensic attribution.")

# ----------------------------------------------------------------------------- references
REFS = {
    "anth_gtg": 'Anthropic Threat Intelligence, "Disrupting the first reported AI-orchestrated cyber espionage campaign," Anthropic, San Francisco, CA, USA, Full Report, Nov. 2025. [Online]. Available: https://assets.anthropic.com/m/ec212e6566a0d47/original/Disrupting-the-first-reported-AI-orchestrated-cyber-espionage-campaign.pdf',
    "mcp": 'Model Context Protocol, "Specification, revision 2025-06-18: Authorization," 2025. [Online]. Available: https://modelcontextprotocol.io/specification/2025-06-18/basic/authorization',
    "a2a": 'Google, "Announcing the Agent2Agent protocol (A2A)," Google for Developers Blog, Apr. 2025. [Online]. Available: https://developers.googleblog.com/en/a2a-a-new-era-of-agent-interoperability/',
    "cyberark": 'CyberArk, "2025 Identity Security Landscape," CyberArk Software Ltd., Newton, MA, USA, Apr. 2025.',
    "dbir": 'Verizon Business, "2025 Data Breach Investigations Report," Verizon, Apr. 2025.',
    "anth_aug": 'Anthropic, "Threat intelligence report: August 2025," Anthropic, Aug. 2025. [Online]. Available: https://www-cdn.anthropic.com/b2a76c6f6992465c09a6f2fce282f6c0cea8c200.pdf',
    "gtig": 'Google Threat Intelligence Group, "GTIG AI threat tracker: Advances in threat actor usage of AI tools," Google Cloud, Nov. 2025. [Online]. Available: https://cloud.google.com/blog/topics/threat-intelligence/threat-actor-usage-of-ai-tools',
    "lyon": 'G. F. Lyon, Nmap Network Scanning: The Official Nmap Project Guide to Network Discovery and Security Scanning. Sunnyvale, CA, USA: Insecure.Com LLC, 2009.',
    "fang": 'R. Fang, R. Bindu, A. Gupta, Q. Zhan, and D. Kang, "LLM agents can autonomously hack websites," arXiv:2402.06664, 2024.',
    "deng": 'G. Deng et al., "PentestGPT: Evaluating and harnessing large language models for automated penetration testing," in Proc. 33rd USENIX Security Symp., Philadelphia, PA, USA, 2024.',
    "greshake": 'K. Greshake, S. Abdelnabi, S. Mishra, C. Endres, T. Holz, and M. Fritz, "Not what you\'ve signed up for: Compromising real-world LLM-integrated applications with indirect prompt injection," in Proc. 16th ACM Workshop on Artificial Intelligence and Security (AISec), 2023, pp. 79-90.',
    "owasp_llm": 'OWASP GenAI Security Project, "OWASP Top 10 for LLM applications 2025," OWASP Foundation, Nov. 2024. [Online]. Available: https://genai.owasp.org/llm-top-10/',
    "nist_ai": 'A. Vassilev et al., "Adversarial machine learning: A taxonomy and terminology of attacks and mitigations," National Institute of Standards and Technology, NIST AI 100-2 E2025, Mar. 2025.',
    "atlas": 'MITRE, "ATLAS: Adversarial threat landscape for artificial-intelligence systems." [Online]. Available: https://atlas.mitre.org',
    "owasp_nhi": 'OWASP Foundation, "OWASP Top 10 non-human identities risks - 2025," 2025. [Online]. Available: https://owasp.org/www-project-non-human-identities-top-10/2025/top-10-2025/',
    "owasp_agentic": 'OWASP GenAI Security Project, "OWASP Top 10 for agentic applications for 2026," OWASP Foundation, Dec. 2025. [Online]. Available: https://genai.owasp.org/resource/owasp-top-10-for-agentic-applications-for-2026/',
    "spiffe": 'SPIFFE Project, "Secure Production Identity Framework for Everyone (SPIFFE) overview," Cloud Native Computing Foundation. [Online]. Available: https://spiffe.io/docs/latest/spiffe-about/overview/',
    "zt": 'S. Rose, O. Borchert, S. Mitchell, and S. Connelly, "Zero trust architecture," National Institute of Standards and Technology, NIST SP 800-207, Aug. 2020.',
    "chan": 'A. Chan et al., "Visibility into AI agents," in Proc. ACM Conf. Fairness, Accountability, and Transparency (FAccT), Rio de Janeiro, Brazil, 2024.',
    "south": 'T. South, S. Marro, T. Hardjono, R. Mahari, C. D. Whitney, D. Greenwood, A. Chan, and A. Pentland, "Authenticated delegation and authorized AI agents," arXiv:2501.09674, Jan. 2025.',
    "saltzer": 'J. H. Saltzer and M. D. Schroeder, "The protection of information in computer systems," Proc. IEEE, vol. 63, no. 9, pp. 1278-1308, Sep. 1975.',
    "hardy": 'N. Hardy, "The confused deputy (or why capabilities might have been invented)," ACM SIGOPS Operating Systems Review, vol. 22, no. 4, pp. 36-38, Oct. 1988.',
    "rfc6749": 'D. Hardt, Ed., "The OAuth 2.0 authorization framework," IETF RFC 6749, Oct. 2012.',
    "rfc8693": 'M. Jones, A. Nadalin, B. Campbell, J. Bradley, and C. Mortimore, "OAuth 2.0 token exchange," IETF RFC 8693, Jan. 2020.',
    "rfc7519": 'M. Jones, J. Bradley, and N. Sakimura, "JSON Web Token (JWT)," IETF RFC 7519, May 2015.',
    "rfc8707": 'B. Campbell, J. Bradley, and H. Tschofenig, "Resource indicators for OAuth 2.0," IETF RFC 8707, Feb. 2020.',
    "macaroons": 'A. Birgisson, J. G. Politz, U. Erlingsson, A. Taly, M. Vrable, and M. Lentczner, "Macaroons: Cookies with contextual caveats for decentralized authorization in the cloud," in Proc. Network and Distributed System Security Symp. (NDSS), San Diego, CA, USA, 2014.',
    "ietf_obo": 'T. S. Senarath and A. Dissanayaka, "OAuth 2.0 extension: On-behalf-of user authorization for AI agents," IETF Internet-Draft draft-oauth-ai-agents-on-behalf-of-user-02, Aug. 2025.',
    "rfc2104": 'H. Krawczyk, M. Bellare, and R. Canetti, "HMAC: Keyed-hashing for message authentication," IETF RFC 2104, Feb. 1997.',
    "denning": 'D. E. Denning, "An intrusion-detection model," IEEE Trans. Software Engineering, vol. SE-13, no. 2, pp. 222-232, Feb. 1987.',
    "roesch": 'M. Roesch, "Snort: Lightweight intrusion detection for networks," in Proc. 13th USENIX Systems Administration Conf. (LISA), Seattle, WA, USA, 1999, pp. 229-238.',
    "paxson": 'V. Paxson, "Bro: A system for detecting network intruders in real-time," Computer Networks, vol. 31, no. 23-24, pp. 2435-2463, Dec. 1999.',
    "nist80094": 'K. Scarfone and P. Mell, "Guide to intrusion detection and prevention systems (IDPS)," National Institute of Standards and Technology, NIST SP 800-94, Feb. 2007.',
    "sommer": 'R. Sommer and V. Paxson, "Outside the closed world: On using machine learning for network intrusion detection," in Proc. IEEE Symp. Security and Privacy, Oakland, CA, USA, 2010, pp. 305-316.',
    "hutchins": 'E. M. Hutchins, M. J. Cloppert, and R. M. Amin, "Intelligence-driven computer network defense informed by analysis of adversary campaigns and intrusion kill chains," in Proc. 6th Int. Conf. Information Warfare and Security (ICIW), 2011, pp. 113-125.',
    "attack": 'B. E. Strom, A. Applebaum, D. P. Miller, K. C. Nickels, A. G. Pennington, and C. B. Thomas, "MITRE ATT&CK: Design and philosophy," The MITRE Corporation, McLean, VA, USA, Tech. Rep. MTR180360, 2018.',
    "nist80086": 'K. Kent, S. Chevalier, T. Grance, and H. Dang, "Guide to integrating forensic techniques into incident response," National Institute of Standards and Technology, NIST SP 800-86, Aug. 2006.',
    "casey": 'E. Casey, Digital Evidence and Computer Crime: Forensic Science, Computers, and the Internet, 3rd ed. Waltham, MA, USA: Academic Press, 2011.',
    "schneier": 'B. Schneier and J. Kelsey, "Secure audit logs to support computer forensics," ACM Trans. Information and System Security, vol. 2, no. 2, pp. 159-176, May 1999.',
    "crosby": 'S. A. Crosby and D. S. Wallach, "Efficient data structures for tamper-evident logging," in Proc. 18th USENIX Security Symp., Montreal, QC, Canada, 2009, pp. 317-334.',
    "nist80061": 'A. Nelson, S. Rekhi, M. Souppaya, and K. Scarfone, "Incident response recommendations and considerations for cybersecurity risk management: A CSF 2.0 community profile," National Institute of Standards and Technology, NIST SP 800-61r3, Apr. 2025.',
}

# ----------------------------------------------------------------------------- body content
blocked_pct = N["misuse_blocked_pct"]
extra_us = round(US["nhi"] - US["baseline"], 1)

BODY = [
    ("h1", "Abstract"),
    ("p", "In mid-September 2025, Anthropic detected and disrupted a cyber-espionage operation that it attributed "
          "with high confidence to a Chinese state-sponsored group designated GTG-1002. An AI coding agent carried out "
          "an estimated 80-90% of the tactical work against roughly thirty organisations [@anth_gtg]. The operators "
          "did not need custom malware: they split the operation into small tasks for sub-agents, each of which "
          "looked legitimate on its own, and connected those agents to ordinary security tools through the Model "
          "Context Protocol. This report studies the campaign from the defender's side. Our argument is that the "
          "missing control was identity, because the agents acted with broad authority that could not be traced to an "
          "accountable person. We reconstruct the campaign from the public report, map each phase to MITRE ATT&CK "
          "tactics and to the identity failure behind it, and design three controls: a non-human identity (NHI) for "
          "every agent, scoped delegation tokens that can only narrow as they pass from agent to sub-agent, and an "
          f"identity-aware intrusion detection layer fed by a hash-chained audit log. A small Python prototype, "
          f"tested on synthetic workloads over {N['runs']} seeded runs, blocked {blocked_pct:.0f}% of misuse calls, "
          f"cut exposed production resources from {B['blast_radius_mean']:.1f} to {N['blast_radius_mean']:.0f} per run, "
          f"removed false alerts ({B['false_alerts_mean']:.2f} to {N['false_alerts_mean']:.0f} per run) and tied every true "
          f"alert to the responsible human, for about {extra_us} µs of extra cost per tool call."),
    ("kw", "Non-human identity; scoped delegation; agentic AI security; intrusion detection; digital forensics; "
           "Model Context Protocol; GTG-1002"),

    ("h1", "1. Introduction"),
    ("p", "AI systems have moved from answering questions to taking actions. Open protocols such as the Model "
          "Context Protocol (MCP) [@mcp] and Agent2Agent (A2A) [@a2a] let an agent call tools, read data and hand "
          "work to other agents. Each of these agents is a non-human identity, and such identities already dominate "
          "enterprise estates: machine identities outnumber human ones by about 82 to 1, and 42% of them hold "
          "privileged or sensitive access [@cyberark]. Stolen or abused credentials remain the most common way into a "
          "network, at 22% of initial access vectors in the 2025 Verizon report [@dbir]. Agents make the problem "
          "sharper because they are created quickly, often share one API key, and delegate tasks to other agents "
          "without any record of who authorised what."),
    ("h2", "1.1 Case background: the GTG-1002 campaign"),
    ("p", "Anthropic's threat intelligence team detected the activity in mid-September 2025 and published its "
          "full report in November 2025 [@anth_gtg]. The campaign targeted about thirty entities, including large "
          "technology companies, financial institutions, chemical manufacturers and government agencies, and the "
          "investigation validated a handful of successful intrusions. Human operators chose the targets and approved "
          "a few escalation points, such as moving from reconnaissance to exploitation, using harvested credentials "
          "and deciding what data to take. The AI performed the remaining 80-90% of the work, at peak rates of several "
          "operations per second that no human team could match. To get past the model's safeguards, the operators "
          "posed as staff of a legitimate security firm doing defensive testing, and broke the work into small tasks "
          "that each looked harmless. The agents drove commodity open-source tools, such as network scanners, "
          "database exploitation frameworks and password crackers, through custom MCP servers. The AI also "
          "overstated its results at times, for example reporting credentials that did not work. Anthropic banned the "
          "accounts, notified affected organisations and authorities, and improved its detection classifiers. The "
          "case goes beyond the company's earlier \"vibe hacking\" report, in which humans still directed most of "
          "the work [@anth_aug], and it sits alongside Google's reports of malware that queries language models at "
          "run time [@gtig]."),
    ("h2", "1.2 Relevance to Unit 4"),
    ("p", "The case touches four Unit 4 topics. Intrusion detection: the campaign was eventually caught because "
          "its tempo and multi-stage pattern stood out, and our prototype adds identity-aware IDS rules for agent "
          "traffic. Digital forensics: investigators had to reconstruct which agent did what on whose behalf, which is "
          "an attribution problem. Security tools: the toolkit was the same class of tools practised in the Kali Linux "
          "and Nmap laboratory sessions [@lyon], with the difference that an agent, not a person, was operating them. "
          "Malware analysis: because no custom malware was used, signature-based analysis had little to inspect, so "
          "detection must rely on behaviour and identity instead."),
    ("h2", "1.3 Objectives"),
    ("bullets", [
        "Analyse GTG-1002 phase by phase and identify the identity failure that each phase relied on.",
        "Design controls based on non-human identity and scoped delegation that feed intrusion detection and forensics.",
        "Build a small prototype and measure prevention, detection, false alerts, attribution and overhead against a shared-key baseline.",
    ]),

    ("h1", "2. Literature Review"),
    ("h2", "2.1 AI-orchestrated attacks"),
    ("p", "Research prototypes showed early that language-model agents can complete offensive tasks. Fang et al. "
          "reported agents that exploited test websites without human help [@fang], and PentestGPT showed that a model "
          "can plan and track multi-step penetration tests [@deng]. Greshake et al. showed that agents reading "
          "untrusted content can be steered by indirect prompt injection [@greshake]. OWASP's Top 10 for LLM "
          "applications lists excessive agency as a key risk [@owasp_llm], NIST's adversarial machine learning "
          "taxonomy covers attacks on and through AI systems [@nist_ai], and MITRE ATLAS catalogues adversary "
          "techniques against AI [@atlas]. Industry threat reports in 2025 then moved from proofs of concept to "
          "real operations [@anth_aug], [@gtig], [@anth_gtg]. Most of this work concentrates on what the model can "
          "do or how the provider can refuse; less of it asks how the surrounding environment authorises each "
          "action an agent takes."),
    ("h2", "2.2 Non-human and agent identity"),
    ("p", "The OWASP Non-Human Identities Top 10 names risks such as improper offboarding, secret leakage, "
          "overprivileged identities and identity reuse [@owasp_nhi]. The OWASP Top 10 for agentic applications adds "
          "agent-specific risks, including identity and privilege abuse and insecure inter-agent communication "
          "[@owasp_agentic]. Workload identity frameworks such as SPIFFE give each service a verifiable identity "
          "[@spiffe], in line with the zero-trust principle that no request is trusted because of where it comes from "
          "[@zt]. Chan et al. propose three visibility measures for AI agents: agent identifiers, real-time monitoring "
          "and activity logs [@chan]. South et al. extend OAuth 2.0 and OpenID Connect so that a human can delegate "
          "limited, auditable authority to an agent [@south]."),
    ("h2", "2.3 Delegation and least privilege"),
    ("p", "Saltzer and Schroeder's principles of least privilege and complete mediation remain the foundation "
          "[@saltzer]. Hardy's confused deputy shows what happens when a program uses its own authority for "
          "someone else's request [@hardy], which is exactly the position of an agent working on a human's behalf. "
          "OAuth 2.0 [@rfc6749] and token exchange [@rfc8693] support delegation, with an actor claim that records "
          "who acts for whom, and JSON Web Tokens carry such claims [@rfc7519]. Resource indicators bind a token to "
          "one service [@rfc8707] and are required by the MCP authorization specification [@mcp]. Macaroons showed "
          "that a bearer credential can be attenuated, so that each holder can only add restrictions [@macaroons]. "
          "An IETF draft now proposes an on-behalf-of flow designed for AI agents [@ietf_obo]."),
    ("h2", "2.4 Intrusion detection and digital forensics"),
    ("p", "Denning's model framed intrusion detection as spotting deviations from normal subject behaviour "
          "[@denning]. Snort [@roesch] and Bro, now Zeek [@paxson], made network IDS practical, and NIST SP 800-94 "
          "describes how detection and prevention systems are deployed [@nist80094]. Sommer and Paxson warn that "
          "anomaly detection struggles with high false-positive rates in real networks [@sommer], a point that "
          "matters for agent traffic. The kill chain [@hutchins] and MITRE ATT&CK [@attack] give a shared vocabulary "
          "for attack stages. On the forensic side, NIST SP 800-86 [@nist80086] and Casey [@casey] stress the "
          "integrity of evidence, and Schneier and Kelsey [@schneier] and Crosby and Wallach [@crosby] show how "
          "logs can be made tamper-evident. NIST SP 800-61r3 places incident response inside overall risk management "
          "[@nist80061]."),
    ("h2", "2.5 Research gap"),
    ("p", "These strands are mostly studied apart. Identity work rarely measures the effect on detection, and IDS "
          "work rarely assumes that each request carries a verifiable delegation chain. As far as we found, no "
          "public study evaluates per-agent identity and scoped delegation as inputs to intrusion detection and "
          "forensic attribution against a scenario shaped like GTG-1002. This report takes a small step toward "
          "closing that gap."),

    ("h1", "3. Methodology"),
    ("p", "The study follows four stages. Each stage produces the input for the next, so the controls and the "
          "experiment can be traced back to evidence from the real case."),
    ("p", "**Stage 1: Case reconstruction.** We used Anthropic's full report as the primary source [@anth_gtg] and "
          "industry reports for context [@anth_aug], [@gtig]. We extracted the six phases, the human approval "
          "gates, the tool categories and the operational tempo. Only public, summary-level information was used, "
          "and no part of the attack was reproduced."),
    ("p", "**Stage 2: Threat and identity-failure mapping.** Each phase was mapped to MITRE ATT&CK tactics "
          "[@attack] and to the identity weakness it depended on, using the OWASP NHI and agentic risk lists "
          "[@owasp_nhi], [@owasp_agentic]. For each weakness we chose a control and noted the Unit 4 area it "
          "supports. Table 1 shows the result."),
    ("table1", None),
    ("p", "**Stage 3: Control design.** Five principles guided the design. (P1) Every agent has its own identity "
          "and a named human owner, so that every action can be traced to a person [@owasp_nhi], [@chan]. (P2) "
          "Authority can only narrow along a delegation chain: a sub-agent can never receive more scope, more "
          "resources, a longer lifetime or more delegation depth than its parent [@macaroons], [@rfc8693], [@south]. "
          "(P3) Every tool call passes through one gateway that checks the token, which is complete mediation "
          "[@saltzer] applied in a zero-trust way [@zt]. (P4) Every decision, allowed or denied, is written to a "
          "hash-chained log [@schneier], [@crosby]. (P5) Detection rules are keyed on the agent identity and the "
          "delegation grant rather than on an IP address or a shared key. We preferred these controls to two "
          "alternatives. Model-side classifiers are valuable, but an organisation does not control them and "
          "attackers can switch models. Network IDS sees mostly encrypted calls to legitimate APIs. Identity is the "
          "one signal the defending organisation can enforce and log itself."),
    ("p", "**Stage 4: Experimental evaluation.** We compared two configurations on the same synthetic workload: a "
          "baseline in which every agent shares one long-lived API key with wildcard scope, which is common in "
          "practice, and the proposed design. Both used identical IDS rules and thresholds, so any difference comes "
          "from identity alone. Synthetic data was chosen for ethical and legal reasons, since no real system is "
          "touched, and because it gives labelled ground truth and repeatable runs. We ran "
          f"{N['runs']} seeds per configuration. The metrics were: the share of misuse calls blocked; blast radius, "
          "meaning distinct production or external resources reached by misuse; false alerts per run; time from "
          "the first misuse event to the first true alert; the share of true alerts that name the accountable human; "
          "gateway cost per call; and whether the audit log detects tampering. Ground-truth labels were used only "
          "for scoring and were never visible to the IDS."),

    ("h1", "4. Implementation Details"),
    ("p", "The prototype is written in Python 3 using only the standard library (hmac, hashlib, json, secrets) plus "
          "matplotlib for the charts, in about 500 lines. Fig. 1 shows the architecture and Table 2 lists the modules. "
          "Tool names are abstract labels and nothing is executed; the prototype only decides whether a call would "
          "be allowed, records the decision and analyses the record."),
    ("fig", "fig_architecture.png", 6.2, "Architecture of the prototype: identity and delegation (top) and the detection and forensics pipeline (bottom)."),
    ("table2", None),
    ("p", "**Tokens.** A token is the Base64 encoding of a JSON claim set followed by an HMAC-SHA256 signature "
          "[@rfc2104]. The claims are the token id, the grant id (shared by the whole delegation tree), the subject "
          "agent, the delegation chain starting with the human, the scopes, the resource prefixes, the issue and "
          "expiry times, the remaining delegation depth and the declared purpose. The delegate() function rejects any "
          "request whose scopes or resources are not a subset of the parent's, issues the child an expiry no later "
          "than the parent's, and reduces the depth by one. Refusals are reported as scope_widening, "
          "resource_widening or depth_exceeded and are themselves logged."),
    ("p", "**Gateway and audit log.** Each tool maps to one required scope and to the ATT&CK tactic it would serve "
          "if misused. The gateway verifies the signature, expiry, revocation, scope and resource prefix, then "
          "appends a record containing the time, agent, chain, grant, tool, tactic, resource, decision and reason. "
          "Each record also stores the SHA-256 hash of the previous record and its own hash, so editing any earlier "
          "record breaks every later link."),
    ("p", "**Identity-aware IDS.** Three simple rules run over the audit stream. R1 (tempo) fires when one agent "
          "makes more than 20 calls in 10 seconds, a rate no human operator sustains. R2 (probing) fires when one "
          "agent is denied three times within 60 seconds. R3 (multi-tactic chain) fires when one delegation grant "
          "touches three or more sensitive tactics within 30 minutes; it is keyed on the grant, so tasks split across "
          "sub-agents are correlated again. Each rule and key has a 15-minute cooldown."),
    ("p", "**Workload.** Four benign agents run for three simulated hours: a CI agent with bursty test runs, a "
          "documentation agent, an operations agent that inventories hosts and queries a reporting database, and a "
          "backup agent that exports production tables to a backup vault. The backup job was included on purpose, "
          "because it looks like collection followed by exfiltration. The misuse case follows only the public shape "
          "of GTG-1002. An external contractor identity owns an orchestrator holding an engagement grant for "
          "inventory and scanning on staging systems only, with a two-hour lifetime and a delegation depth of one. "
          "The orchestrator tries to hand five sub-agents wider tasks across discovery, credential access, lateral "
          "movement, collection and exfiltration. Phases are separated by human approval gaps of 2-20 minutes, and "
          "the tempo varies from 0.3 to 4 calls per second between runs."),

    ("h1", "5. Results and Discussion"),
    ("p", "Table 3 summarises the 50-run comparison. Fig. 2 breaks the allowed misuse calls down by tactic, Fig. 3 "
          "compares the main outcomes, and Fig. 4 shows the forensic timeline rebuilt from the audit log for one run."),
    ("table3", None),
    ("fig", "fig_tactics.png", 5.6, "Mean misuse calls allowed per run, by ATT&CK tactic (50 runs)."),
    ("p", f"**Prevention and blast radius.** Scoped delegation blocked {blocked_pct:.0f}% of misuse calls. Every "
          "attempt to hand a sub-agent credential access, remote execution, database export or outbound transfer "
          "was refused at delegation time, and the refusal itself became evidence. The calls that still succeeded "
          f"(about {N['misuse_allowed_mean']:.0f} per run) were discovery calls on staging hosts that the engagement "
          "had actually authorised. Scoping therefore does not stop misuse of authority that was legitimately granted, "
          "so detection is still needed. Production and external resources reached by misuse fell from "
          f"{B['blast_radius_mean']:.1f} per run to zero, and no benign call was denied in either configuration."),
    ("fig", "fig_outcomes.png", 6.2, "Main outcomes: misuse blocked, resources exposed, false alerts and attribution (means over 50 runs)."),
    ("p", f"**Detection.** Both configurations detected the misuse in every run, and the baseline was fast too "
          f"(median {B['ttd_median_s']} s compared with {N['ttd_median_s']} s), because high tempo is visible even "
          "on a shared key. The difference lies in what an alert means. Baseline alerts name only the shared key, so a "
          "responder can contain the attack only by revoking that key and stopping all four legitimate agents along "
          f"with it. Mixing several agents' traffic on one key also produced {B['false_alerts_mean']:.2f} false "
          "alerts per run: the inventory sweeps and the backup job together looked like a multi-tactic chain. "
          f"With per-agent identities there were no false alerts, and all {N['true_alerts_mean']:.1f} true alerts "
          "per run named the agent, its orchestrator and the human owner, so a single grant could be revoked without "
          "affecting anything else. The probing rule (R2), which only works when denials are attributed to a specific "
          "agent, raised the first alert in most runs."),
    ("fig", "fig_timeline.png", 6.2, "Forensic timeline rebuilt from the audit log for one run (seed 7): allowed calls, gateway denials and IDS alerts per identity."),
    ("p", f"**Forensics and cost.** Every misuse record was attributable to the contractor identity "
          f"({N['record_attribution_pct']:.0f}% compared with {B['record_attribution_pct']:.0f}%), so the timeline in "
          "Fig. 4 can be read directly as \"who, on whose behalf, did what, when\". In the tamper test, changing one "
          f"denied record to allowed in a {TT['records']}-record log was detected at exactly record {TT['detected_at_index']}. "
          f"Verifying a token added about {extra_us} µs per call ({US['nhi']} µs compared with {US['baseline']} µs for "
          "the shared key, in pure Python), which is negligible next to the latency of a model call."),
    ("p", "**What this means for GTG-1002.** In the real campaign the agents ran on the attackers' own "
          "infrastructure, so a victim cannot force the attacker to use scoped tokens. The controls therefore apply in "
          "three places. AI and tool platforms can require a verified, accountable identity and an explicit, scoped "
          "grant before an agent is used for security testing, which turns a role-play claim into a decision that a "
          "named person is responsible for. Organisations running their own agents can stop those agents from being "
          "hijacked or over-trusted. Target organisations can issue short-lived, narrowly scoped credentials to "
          "machines, so that credentials harvested during an intrusion are worth far less. In all three cases the "
          "same identity data that limits damage also makes detection more precise and attribution possible."),
    ("p", "**Limitations.** The data is synthetic and the thresholds were set by hand, so the absolute numbers "
          "should not be generalised; real deployments would need tuning against real logs, given the known "
          "difficulty of anomaly detection [@sommer]. The prototype uses a single signing key and does not cover "
          "delegation across organisations, token theft from memory, or an attacker who stays slowly inside a "
          "legitimate scope."),

    ("h1", "6. Conclusion"),
    ("p", "GTG-1002 showed that an AI agent can carry out most of an intrusion using ordinary tools, with humans "
          "stepping in only at a few decision points. The most important lesson for defenders is not about any single "
          "exploit; it is that the agents acted with authority that nobody had scoped and nobody could trace. This "
          "report mapped each phase of the campaign to the identity failure it relied on and tested three controls: a "
          "non-human identity per agent, delegation that can only narrow, and identity-aware detection backed by a "
          f"tamper-evident log. In our prototype these controls blocked {blocked_pct:.0f}% of misuse calls, reduced the "
          "exposed production surface to zero, removed false alerts and attributed every alert to an accountable "
          "human, at negligible cost. Future work includes testing on real agent logs, learning thresholds per "
          "agent, cross-organisation delegation over A2A, and platform-level verification of agent identity before "
          "security-sensitive use."),
]

TABLE1 = (["GTG-1002 phase", "ATT&CK tactic", "Identity failure", "Control in this study", "Unit 4 link"], [
    ["1. Initialisation and target selection (security-firm persona)", "Resource development", "Claimed purpose not bound to an accountable identity", "Grant issued by a named human; purpose stored in the token", "Forensics"],
    ["2. Reconnaissance and attack-surface mapping", "Discovery", "One broad credential across all systems; machine tempo", "Resource-prefix scope; R1 tempo; R2 probing", "IDS; tools (scanners)"],
    ["3. Vulnerability discovery and validation", "Initial access, execution", "Sub-tasks look legitimate in isolation", "Per-sub-agent identity; R3 keyed on the whole grant", "IDS"],
    ["4. Credential harvesting and lateral movement", "Credential access, lateral movement", "Harvested secrets usable by any agent", "Scopes that cannot be widened; short lifetimes", "IDS; forensics"],
    ["5. Data collection and extraction", "Collection, exfiltration", "No limit on blast radius", "Resource scoping; refusals logged", "Forensics"],
    ["6. Documentation and handoff", "Persistence (access handed on)", "Access handed to other teams with no expiry", "TTL, depth limit, revocation; hash-chained log", "Forensics"],
], [1.45, 1.0, 1.4, 1.5, 0.85], "Mapping of GTG-1002 phases to ATT&CK tactics, identity failures and controls (from [@anth_gtg], [@attack]).")

TABLE2 = (["Module", "Responsibility"], [
    ["identity.py", "Identity registry; DelegationAuthority issues, verifies and attenuates HMAC-signed tokens; SharedKeyAuthority for the baseline"],
    ["gateway.py", "Single enforcement point: verify token, check scope and resource, record the decision"],
    ["audit.py", "Append-only SHA-256 hash-chained log with verify() and timeline()"],
    ["ids.py", "Identity-aware rules R1 tempo, R2 probing, R3 multi-tactic chain"],
    ["workload.py", "Synthetic benign agents and the abstract misuse delegation tree"],
    ["run_experiment.py, make_figures.py", "50-seed comparison, metrics, tamper test, overhead benchmark, charts"],
], [1.9, 4.3], "Prototype modules (implementation/nhi_lab).")

TABLE3 = (["Metric (mean of 50 runs unless stated)", "Baseline: shared key", "NHI + scoped delegation"], [
    ["Misuse calls attempted per run", f"{B['misuse_calls_mean']:.1f}", f"{N['misuse_calls_mean']:.1f}"],
    ["Misuse calls blocked", f"{B['misuse_blocked_pct']:.0f}%", f"{N['misuse_blocked_pct']:.1f}%"],
    ["Production/external resources exposed per run", f"{B['blast_radius_mean']:.1f}", f"{N['blast_radius_mean']:.0f}"],
    ["Benign calls wrongly denied (all runs)", f"{B['benign_denied_total']}", f"{N['benign_denied_total']}"],
    ["Runs in which misuse was detected", f"{B['detection_rate_pct']:.0f}%", f"{N['detection_rate_pct']:.0f}%"],
    ["Median time to first true alert", f"{B['ttd_median_s']} s", f"{N['ttd_median_s']} s"],
    ["False alerts per run", f"{B['false_alerts_mean']:.2f}", f"{N['false_alerts_mean']:.0f}"],
    ["True alerts naming the accountable human", f"{B['alert_attribution_pct']:.0f}%", f"{N['alert_attribution_pct']:.0f}%"],
    ["Misuse audit records attributable", f"{B['record_attribution_pct']:.0f}%", f"{N['record_attribution_pct']:.0f}%"],
    ["Gateway cost per call", f"{US['baseline']} µs", f"{US['nhi']} µs"],
], [3.3, 1.45, 1.45], "Results of the 50-run comparison.")

# ----------------------------------------------------------------------------- helpers
order = []


def cite(text):
    def num(key):
        if key not in REFS:
            raise KeyError(key)
        if key not in order:
            order.append(key)
        return order.index(key) + 1

    def group(m):
        nums = sorted(num(k) for k in re.findall(r"@(\w+)", m.group(0)))
        return "[" + ", ".join(map(str, nums)) + "]"
    # a run like "[@a], [@b]" becomes one sorted group "[3, 7]"
    return re.sub(r"\[@\w+\](?:, \[@\w+\])*", group, text)


def set_font(run, size=12, bold=None, italic=None, color=BLACK, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.0, keep_next=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next


def add_rich(p, text, size=12, color=BLACK):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            set_font(p.add_run(part), size=size, bold=(i % 2 == 1), color=color)


def heading(doc, text, level):
    p = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    fmt_para(p, WD_ALIGN_PARAGRAPH.LEFT, before=12 if level == 1 else 8, after=4, keep_next=True)
    set_font(p.add_run(text), size=14 if level == 1 else 12, bold=True, color=BLUE)


def caption(doc, text, before=2, after=10):
    p = doc.add_paragraph()
    fmt_para(p, WD_ALIGN_PARAGRAPH.CENTER, before=before, after=after)
    set_font(p.add_run(cite(text)), size=11, italic=True)
    return p


def cell_shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc, spec, number):
    headers, rows, widths, cap = spec
    caption(doc, f"Table {number}. {cap}", before=6, after=4).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = doc.styles["Table Grid"] if "Table Grid" in [s.name for s in doc.styles] else None
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "1F3864")
        borders.append(el)
    tblpr.append(borders)
    for r_i, values in enumerate([headers] + rows):
        row = t.rows[r_i]
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); trpr.append(cant)
        for c_i, val in enumerate(values):
            cell = row.cells[c_i]
            cell.width = Inches(widths[c_i])
            p = cell.paragraphs[0]
            fmt_para(p, WD_ALIGN_PARAGRAPH.LEFT, after=0)
            set_font(p.add_run(cite(val)), size=10.5, bold=(r_i == 0), color=BLUE if r_i == 0 else BLACK)
            if r_i == 0:
                cell_shade(cell, "E8EEF8")
    spacer = doc.add_paragraph(); fmt_para(spacer, after=2)


def add_figure(doc, name, width, cap, number):
    p = doc.add_paragraph()
    fmt_para(p, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, keep_next=True)
    p.add_run().add_picture(str(FIGS / name), width=Inches(width))
    caption(doc, f"Fig. {number}. {cap}")


def set_cell_text(cell, text, font=FONT, size=None):
    p = cell.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    set_font(runs[0], size=size or runs[0].font.size.pt, color=None, name=font)


def replace_para(p, parts, size, color=None):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, bold in parts:
        set_font(p.add_run(text), size=size, bold=bold, color=color)


# ----------------------------------------------------------------------------- build
def edit_cover(doc):
    top = doc.tables[0]
    set_cell_text(top.cell(0, 0), "Academic Year -: 26 – 27", size=10)
    set_cell_text(top.cell(0, 2), "Date: 10/10/2026", size=10)

    paras = doc.paragraphs
    for p in paras:
        txt = p.text.strip()
        if txt in ("SYMBIOSIS INSTITUTE", "OF TECHNOLOGY (SIT)"):
            for r in p.runs:
                set_font(r, size=25, bold=True, color=None, name="Optima")
        elif txt.startswith("Constituent of"):
            for r in p.runs:
                set_font(r, size=7.5, bold=True, color=BLACK)
        elif txt == "DevOps":
            replace_para(p, [("Cyber Security", True)], 15, BLUE)
            p.paragraph_format.space_after = Pt(18)
        elif txt.startswith("CA2 - DevOps"):
            replace_para(p, [("CA-2 (Part 2) – Case Study Report (Unit 4)", True)], 14, BLACK)
            r = p.add_run(); r.add_break()
            set_font(p.add_run("Group Number: Group – 04"), size=13, bold=False, color=BLACK)
            p.paragraph_format.space_after = Pt(22)
        elif txt.startswith("Title:"):
            replace_para(p, [("Title: ", True), (f"“{TITLE}”", False)], 13, BLACK)
            p.paragraph_format.space_after = Pt(18)
        elif txt.startswith("Problem Statement:"):
            replace_para(p, [("Problem Statement: ", True), (PROBLEM, False)], 11, BLACK)
            p.paragraph_format.space_after = Pt(24)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # first spacer after the top table
    paras[0].paragraph_format.space_before = Pt(26)

    members = doc.tables[1]
    template = members.rows[-1]._tr
    while len(members.rows) < 1 + len(MEMBERS):
        template.addnext(copy.deepcopy(template))
    for i, (name, prn) in enumerate(MEMBERS, start=1):
        for c, val in enumerate((name, prn, "CSE-C")):
            set_cell_text(members.cell(i, c), val, size=11)
    for c in range(3):
        set_cell_text(members.cell(0, c), members.cell(0, c).text, size=11)


def edit_header_footer(doc):
    sec = doc.sections[-1]
    hp = sec.header.paragraphs[0]
    replace_para(hp, [(f"CA-2 – Cyber Security · Group 04 Case Study Report", False)], 9, BLUE)
    for p in list(sec.header.paragraphs) + list(sec.footer.paragraphs):
        for r in p.runs:
            set_font(r, size=9, color=BLUE)
        for b in p._p.iter(qn("w:top")):
            b.set(qn("w:color"), "1F3864")
        for b in p._p.iter(qn("w:bottom")):
            b.set(qn("w:color"), "1F3864")


def title_block(doc):
    body_start = doc.paragraphs[-1]   # empty paragraph after the cover section break
    fmt_para(body_start, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    set_font(body_start.add_run(TITLE), size=16, bold=True, color=BLUE)
    p = doc.add_paragraph(); fmt_para(p, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    set_font(p.add_run(", ".join(f"{n} ({prn})" for n, prn in MEMBERS)), size=12)
    p = doc.add_paragraph(); fmt_para(p, WD_ALIGN_PARAGRAPH.CENTER, after=8)
    set_font(p.add_run("Group 04 · B.Tech CSE-C, Semester VII · Department of Computer Engineering, "
                       "Symbiosis Institute of Technology, Pune"), size=11, italic=True)


def main():
    # figures are stored as ("fig", name, width, caption) 4-tuples; normalise them
    global BODY
    BODY = [(b[0], b[1:] if b[0] == "fig" else b[1]) for b in BODY]
    doc = Document(str(HERE / "ForgeFlow_Deployment_Report.docx"))
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    edit_cover(doc)
    edit_header_footer(doc)
    title_block(doc)

    fig_no = tab_no = 0
    for kind, content in BODY:
        if kind in ("h1", "h2"):
            heading(doc, content, 1 if kind == "h1" else 2)
        elif kind == "p":
            p = doc.add_paragraph(); fmt_para(p)
            add_rich(p, cite(content))
        elif kind == "kw":
            p = doc.add_paragraph(); fmt_para(p, after=4)
            set_font(p.add_run("Keywords: "), bold=True, color=BLUE)
            set_font(p.add_run(content), italic=True)
        elif kind == "bullets":
            for item in content:
                p = doc.add_paragraph(); fmt_para(p, after=3)
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.2)
                set_font(p.add_run("•  " + cite(item)))
        elif kind == "fig":
            fig_no += 1
            add_figure(doc, *content, fig_no)
        elif kind.startswith("table"):
            tab_no += 1
            add_table(doc, {"table1": TABLE1, "table2": TABLE2, "table3": TABLE3}[kind], tab_no)

    heading(doc, "References", 1)
    unused = [k for k in REFS if k not in order]
    if unused:
        raise SystemExit(f"uncited references: {unused}")
    for i, key in enumerate(order, start=1):
        p = doc.add_paragraph(); fmt_para(p, WD_ALIGN_PARAGRAPH.LEFT, after=3)
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        set_font(p.add_run(f"[{i}]\t{REFS[key]}"), size=11)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))

    out = HERE / "G04.docx"
    doc.core_properties.title = TITLE
    doc.core_properties.author = "Group 04 – Raghav Sonchhatra, Sanidhya Awasthi, Faheemuddin Sayyed"
    doc.save(str(out))
    print(f"wrote {out.name}: {len(order)} references, {fig_no} figures, {tab_no} tables")


if __name__ == "__main__":
    main()
