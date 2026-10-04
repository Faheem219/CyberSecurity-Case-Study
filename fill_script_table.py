"""Fill the word-count table in G04_Presentation_Script.md (placeholders REPORT_n, SCRIPT_n, TIME_n, TARGET_LEN)."""
import re
from docx import Document

WPM = 140
OWNER = {"Abstract": 1, "1. Introduction": 1, "1.1 Case background: the GTG-1002 campaign": 1,
         "1.2 Relevance to Unit 4": 1, "1.3 Objectives": 2, "2. Literature Review": 1,
         "2.1 AI-orchestrated attacks": 1, "2.2 Non-human and agent identity": 1,
         "2.3 Delegation and least privilege": 2, "2.4 Intrusion detection and digital forensics": 2,
         "2.5 Research gap": 2, "3. Methodology": 2, "4. Implementation Details": 2,
         "5. Results and Discussion": 3, "6. Conclusion": 3, "References": 0}

doc, sec, started, report = Document("G04.docx"), None, False, {1: 0, 2: 0, 3: 0}
for p in doc.paragraphs:
    t = p.text.strip()
    if not t or (not started and t != "Abstract"):
        continue
    started = True
    if p.style.name.startswith("Heading"):
        sec = t
        continue
    who = OWNER[sec]
    if sec == "4. Implementation Details" and t.startswith(("Identity-aware IDS.", "Workload.")):
        who = 3
    if who:
        report[who] += len(t.split())

md = open("G04_Presentation_Script.md").read()
parts = re.split(r"\n## Part \d", md)[1:]
script = {i + 1: len(re.sub(r"###[^\n]*", "", p.split("\n", 1)[1]).replace("---", "").split()) for i, p in enumerate(parts)}
tr, ts = sum(report.values()), sum(script.values())
for i in (1, 2, 3):
    secs = script[i] / WPM * 60
    md = md.replace(f"REPORT_{i}", f"{report[i]} ({100 * report[i] / tr:.0f}%)")
    md = md.replace(f"SCRIPT_{i}", f"{script[i]} ({100 * script[i] / ts:.0f}%)")
    md = md.replace(f"TIME_{i}", f"~{int(secs // 60)} min {int(secs % 60):02d} s")
md = md.replace("TARGET_LEN", f"about {ts / WPM:.1f} minutes in total ({ts} words at a normal speaking pace of ~{WPM} words per minute)")
open("G04_Presentation_Script.md", "w").write(md)
print(report, script, ts, round(ts / WPM, 1))
